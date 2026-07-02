import os
import sys
import subprocess

# --- Automated Dependency Bootstrapper ---
REQUIRED_PACKAGES = ["Flask", "python-docx", "openpyxl", "werkzeug"]

def bootstrap_dependencies():
    """Checks and automatically installs any missing Python packages in the environment."""
    for package in REQUIRED_PACKAGES:
        try:
            # Map programmatic import names to package installation requirements
            if package == "python-docx":
                import docx
            else:
                __import__(package.lower())
        except ImportError:
            print(f"[*] Package '{package}' is missing. Bootstrapping installation...")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", package])
                print(f"[+] Successfully installed '{package}'")
            except Exception as e:
                print(f"[-] Critical Error: Failed to bootstrap package '{package}'. Reason: {e}")
                sys.exit(1)

# Execute the self-installation bootstrap sequence prior to importing library components
bootstrap_dependencies()

import re
import io
import datetime
from flask import Flask, request, jsonify, send_file
from werkzeug.utils import secure_filename
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
from docx.oxml import CT_Tbl, OxmlElement
from docx.oxml.ns import qn

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024  # 32MB Capacity
app.config['UPLOAD_FOLDER'] = '/tmp' if os.name != 'nt' else '.'

TEMPLATE_FILENAME = "new_SRA_report_template.docx"

def normalize_domain_key(text):
    if not text: return "compliance"
    t = str(text).strip().lower().replace(" ", "").replace(",", "").replace("-", "")
    if "managementrespons" in t: return "management"
    if "itsecuritypolic" in t: return "policy"
    if "humanresource" in t: return "human"
    if "assetmanag" in t: return "asset"
    if "accesscontrol" in t: return "access"
    if "cryptograph" in t: return "crypto"
    if "physical" in t: return "physical"
    if "operation" in t: return "operations"
    if "communication" in t: return "communications"
    if "systemacquisition" in t or "developmentandmaintenance" in t: return "development"
    if "outsourcing" in t: return "outsourcing"
    if "incident" in t: return "incident"
    if "businesscontinuity" in t or "aspectsofbc" in t: return "continuity"
    if "compliance" in t: return "compliance"
    return "compliance"

def apply_text_styling_and_borders(cell, text, is_header=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    if is_header:
        run.bold = True
        
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def insert_paragraph_after(paragraph, text, style=None):
    doc_instance = paragraph.part.document
    new_p = doc_instance.add_paragraph(style=style)
    new_p.text = text
    p_element = paragraph._p
    p_element.addnext(new_p._p)
    return new_p

def parse_excel_followup_source(xlsx_path, sys_abbr, global_metrics, domain_matrix, prefix_filter):
    groups = {}
    wb = load_workbook(xlsx_path, data_only=True)
    sheet = wb.active
    
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if row_idx < 2 or row[2] is None: continue
        obs_num = str(row[2]).strip().upper()
        
        if not obs_num.startswith(prefix_filter.upper()): continue
        
        system_affected = sys_abbr
        raw_ip = str(row[3]).strip() if row[3] is not None else ""
        raw_port = str(row[4]).strip() if (len(row) > 4 and row[4] is not None) else ""
        
        if raw_ip and raw_port and ":" not in raw_ip and str(raw_port) not in raw_ip:
            asset_location = f"{raw_ip}:{raw_port}"
        else:
            asset_location = raw_ip if raw_ip else "System Codebase" if prefix_filter == "C" else sys_abbr

        vuln_name = str(row[8]).strip() if row[8] is not None else ""    
        threat_desc = str(row[9]).strip() if row[9] is not None else ""  
        reremediation = str(row[10]).strip() if row[10] is not None else "" 
        severity = str(row[11]).strip() if row[11] is not None else "Low" 
        raw_domain = str(row[5]).strip() if (len(row) > 5 and row[5] is not None) else "Compliance"
        
        sev_title = severity.strip().title()
        if "aoi" in severity.lower(): sev_title = "AOI"
        sev_key = "aoi" if sev_title == "AOI" else sev_title.lower()

        if sev_title == "AOI":
            combined_risk_rating = "AOI"
        else:
            combined_risk_rating = f"{sev_title}\n[{sev_title},{sev_title}]"

        unique_hash_key = (vuln_name.lower().strip(), threat_desc.lower().strip())

        if unique_hash_key not in groups:
            groups[unique_hash_key] = {
                "sys": system_affected,
                "locations": [asset_location],
                "name": vuln_name,
                "threat": threat_desc,
                "rating": combined_risk_rating,
                "fix": reremediation,
                "sev_key": sev_key,
                "domain": raw_domain
            }
        else:
            if asset_location not in groups[unique_hash_key]["locations"]:
                groups[unique_hash_key]["locations"].append(asset_location)

        if sev_key in global_metrics:
            global_metrics[sev_key] += 1
        
        dom_mapped_key = normalize_domain_key(raw_domain)
        if sev_key in domain_matrix[dom_mapped_key]:
            domain_matrix[dom_mapped_key][sev_key] += 1
            
    return [[item["sys"], ", ".join(item["locations"]), item["name"], item["threat"], item["rating"], item["fix"]] for item in groups.values()]

def build_total_vulnerabilities_string(high, medium, low, aoi):
    parts = []
    if high > 0: parts.append(f"{high} high")
    if medium > 0: parts.append(f"{medium} medium")
    if low > 0: parts.append(f"{low} low risk")
    if aoi > 0: parts.append(f"{aoi} AOI")
    
    if not parts: return "no vulnerability"
    if len(parts) == 1: return parts[0]
    if len(parts) == 2: return f"{parts[0]} and {parts[1]}"
    return ", ".join(parts[:-1]) + f", and {parts[-1]}"

@app.route('/api/generate-sra', methods=['POST'])
def handle_sra_generation_pipeline():
    # Enforce Core 3-Source Input Payload Checkpoints
    if 'file_wab' not in request.files or 'file_plan' not in request.files or 'file_codereview' not in request.files:
        return jsonify({"error": "Missing mandatory assets. WAB (.docx), Follow-Up Plan (.xlsx), and Code Review (.xlsx) are all required."}), 400
        
    wab_file = request.files['file_wab']
    plan_file = request.files['file_plan']
    code_file = request.files['file_codereview']

    company_name = request.form.get("company_name", "SunnyVision Limited")
    company_abbr = request.form.get("company_abbr", "SV")
    sys_name = request.form.get("system_name", "Digital Works Supervision System")
    sys_abbr = request.form.get("system_abbr", "DWSS")
    testee_name = request.form.get("testee_name", "Architectural Services Department")
    testee_abbr = request.form.get("testee_abbr", "ArchSD")
    
    today = datetime.date.today()
    current_date_text = today.strftime("%B %Y")
    current_date_numeric = today.strftime("%d-%m-%Y")

    wab_path = os.path.join(app.config['UPLOAD_FOLDER'], secure_filename(wab_file.filename))
    plan_path = os.path.join(app.config['UPLOAD_FOLDER'], secure_filename(plan_file.filename))
    code_path = os.path.join(app.config['UPLOAD_FOLDER'], secure_filename(code_file.filename))
    output_path = os.path.join(app.config['UPLOAD_FOLDER'], f"SRA_Report_{sys_abbr}_{today.strftime('%Y%m%d')}.docx")

    wab_file.save(wab_path)
    plan_file.save(plan_path)
    code_file.save(code_path)

    if not os.path.exists(TEMPLATE_FILENAME):
        return jsonify({"error": f"Base SRA Template layout file target reference '{TEMPLATE_FILENAME}' not found in runtime root tree path."}), 500

    try:
        global_metrics = {"high": 0, "medium": 0, "low": 0, "aoi": 0}
        domain_matrix = {k: {"high": 0, "medium": 0, "low": 0, "aoi": 0} for k in [
            "management", "policy", "human", "asset", "access", "crypto", "physical", 
            "operations", "communications", "development", "outsourcing", "incident", "continuity", "compliance"
        ]}

        v_vulnerabilities = parse_excel_followup_source(plan_path, sys_abbr, global_metrics, domain_matrix, "V")
        a_vulnerabilities = parse_excel_followup_source(plan_path, sys_abbr, global_metrics, domain_matrix, "A")
        c_vulnerabilities = parse_excel_followup_source(code_path, sys_abbr, global_metrics, domain_matrix, "C")

        # Dynamic State Machine for WAB Extraction Boundaries
        system_description_text = ""
        security_requirements = []
        assessment_scope_text = []
        assessment_objective_text = []
        
        wab_doc = Document(wab_path)
        capture_desc, capture_req = False, False
        capture_scope, capture_objective = False, False

        for p in wab_doc.paragraphs:
            txt = p.text.strip()
            if not txt: continue
            
            normalized_txt = txt.replace('\xa0', ' ').strip()
            lower_clean = normalized_txt.lower().replace(" ", "").replace(",", "").replace(":", "")
            
            # Context Scope 1: Current Environment Description Section
            if "4.1.1" in normalized_txt and "Current Environment" in normalized_txt:
                capture_desc = True
                continue
            elif "4.1.2" in normalized_txt or "Project Management" in normalized_txt:
                capture_desc = False

            if capture_desc and len(normalized_txt) > 2:
                if not (normalized_txt.startswith("4.1") or normalized_txt.startswith("4.1.1")):
                    if "USER REQUIREMENTS" not in normalized_txt and "Current Environment Description" not in normalized_txt:
                        system_description_text = normalized_txt if not system_description_text else system_description_text + "\n" + normalized_txt

            # Context Scope 2: Assessment Scope (Section 1.1)
            if "scopeoftheservices" in lower_clean:
                capture_scope = True
                continue
            elif capture_scope and ("background" in lower_clean or "projectobjectives" in lower_clean or "1.2" in lower_clean):
                capture_scope = False

            if capture_scope:
                if not lower_clean.startswith("scope") and "categoryb" not in lower_clean:
                    assessment_scope_text.append(normalized_txt)

            # Context Scope 3: Assessment Objectives (Section 3)
            if "projectobjectives" in lower_clean:
                capture_objective = True
                continue
            elif capture_objective and ("projectrequirements" in lower_clean or "userrequirements" in lower_clean or "4." in lower_clean):
                capture_objective = False

            if capture_objective:
                if not lower_clean.startswith("projectobjective"):
                    assessment_objective_text.append(normalized_txt)

            # Context Scope 4: Security Baseline Extraction Boundary Rules
            if "form" in lower_clean and "securitybaseline" in lower_clean and "assessmentandaudit" in lower_clean:
                capture_req = True
                continue
            elif capture_req and "gather" in lower_clean:
                capture_req = False

            if capture_req:
                if "baseline" not in lower_clean and "form" not in lower_clean:
                    cleaned_bullet = re.sub(r'^\([ivxLCDM]+\)\s*', '', normalized_txt).strip()
                    if cleaned_bullet: security_requirements.append(cleaned_bullet)

        total_vuln_str = build_total_vulnerabilities_string(global_metrics["high"], global_metrics["medium"], global_metrics["low"], global_metrics["aoi"])
        
        replacements = [
            ("*company name* Limited", company_name),
            ("*company name*", company_name),
            ("company name limited", company_name),
            ("company name", company_name),
            ("information system name (information system abbreviation)", f"{sys_name} ({sys_abbr})"),
            ("*(INFORMATION SYSTEM ABBREVIATION)*", sys_abbr),
            ("(INFORMATION SYSTEM ABBREVIATION)", sys_abbr),
            ("INFORMATION SYSTEM ABBREVIATION", sys_abbr),
            ("information system name", sys_name),
            ("information system abbreviation", sys_abbr),
            ("company name (company abbreviation)", f"{company_name} ({company_abbr})"),
            ("company abbreviation", company_abbr),
            ("COMPANY ABBREVIATION", company_abbr),
            ("testee name (testee abbreviation)", f"{testee_name} ({testee_abbr})"),
            ("testee abbreviation", testee_abbr),
            ("TESTEE ABBREVIATION", testee_abbr),
            ("testee name", testee_name),
            ("*testee abbreviation*", testee_abbr),
            ("May 2026", current_date_text),
            ("June 2026", current_date_text),
            ("*grand total items*", total_vuln_str),
            ("*grand total items items*", total_vuln_str)
        ]

        doc = Document(TEMPLATE_FILENAME)
        
        def run_replace_mechanism(p, lookup_tuples):
            full_text = "".join(run.text for run in p.runs)
            if not full_text.strip(): return
            if re.match(r'^\d+(\.\d+)*\s*', full_text.strip()) or p.style.name.startswith("Heading"): return
            
            modified = False
            for k, val in lookup_tuples:
                if k.lower() in full_text.lower() or k.strip("*").strip().lower() in full_text.lower():
                    pattern = re.compile(re.escape(k), re.IGNORECASE)
                    full_text = pattern.sub(val, full_text)
                    modified = True
            if modified or "*" in full_text:
                full_text = full_text.replace("*", "").strip()
                if p.runs:
                    p.runs[0].text = full_text
                    for r in p.runs[1:]: r.text = ""
                else: p.text = full_text

        # Overwrite Global Document Section Headers
        for section in doc.sections:
            for header_p in section.header.paragraphs: run_replace_mechanism(header_p, replacements)
            for footer_p in section.footer.paragraphs: run_replace_mechanism(footer_p, replacements)

        body_elements = doc.element.body
        tbl_index_map = {child: idx for idx, child in enumerate(body_elements) if isinstance(child, CT_Tbl)}
        has_populated_section_8 = False

        for p_idx, p in enumerate(list(doc.paragraphs)):
            lower_normalized = p.text.lower().replace(" ", "").replace(",", "")
            
            if p_idx < 15 and ("may 2026" in p.text.lower() or "june 2026" in p.text.lower()):
                p.text = current_date_text
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
                continue

            if "8.securityrequirements" in lower_normalized or p.style.name.startswith("Heading"): continue

            # Dynamic Scope Stream Injection Points
            if "insert all under brief of work assignment under category b, 1. Scope of the services, 1.1".replace(" ", "") in lower_normalized or "copy from WAB, 1.1 Scope".replace(" ", "") in lower_normalized:
                p.text = ""
                current_anchor = p
                if assessment_scope_text:
                    for text_line in assessment_scope_text:
                        current_anchor = insert_paragraph_after(current_anchor, text_line, style='List Bullet')
                else:
                    p.text = "The scope of the services covers the security areas specified in S17 and G3 guidelines."

            # Dynamic Objectives Stream Injection Points
            elif "copy from work assignment under category b, all under 3. Project objectives".replace(" ", "") in lower_normalized or "copy from WAB, 3. Project Objectives".replace(" ", "") in lower_normalized:
                p.text = ""
                current_anchor = p
                if assessment_objective_text:
                    for text_line in assessment_objective_text:
                        current_anchor = insert_paragraph_after(current_anchor, text_line, style='List Bullet')
                else:
                    p.text = "To evaluate security risks and verify regulatory baseline compliance."

            elif "copyfromwabsection4.4" in lower_normalized or "copyfromwabsection4.1.3" in lower_normalized:
                if not has_populated_section_8:
                    has_populated_section_8 = True
                    p.text = ""
                    current_anchor = p
                    target_list = security_requirements if security_requirements else [
                        "Baseline IT Security Policy;", "IT Security Guidelines;", "Practice Guide for Security Risk Assessment & Audit;",
                        "Practice Guide for Information Security Incident Handling;", "Practice Guide for Penetration Testing;", 
                        "Practice Guide for Website and Web Application Security;", "Practice Guide for Cloud Computing Security;", 
                        "Practice Guide for Mobile Security;", "Practice Guide for Internet of Things Security"
                    ]
                    for item in target_list:
                        current_anchor = insert_paragraph_after(current_anchor, item, style='List Bullet')
                else:
                    p.text = ""
            elif "copy from WAB, 4.1.1 Current Environment Description" in p.text or "copy from WAB, 9.1 System Description" in p.text:
                p.text = system_description_text
            elif "copy from follow up plan" in p.text:
                if "v numbered rows" in p.text.lower(): target_source = v_vulnerabilities
                elif "a numbered rows" in p.text.lower(): target_source = a_vulnerabilities
                else: target_source = c_vulnerabilities
                
                current_node = p._p.getnext()
                while current_node is not None and not isinstance(current_node, CT_Tbl):
                    current_node = current_node.getnext()
                    
                if current_node is not None and current_node in tbl_index_map:
                    target_table = doc.tables[list(tbl_index_map.keys()).index(current_node)]
                    for row_item in target_source:
                        new_row = target_table.add_row()
                        for cell_idx, val in enumerate(row_item):
                            if cell_idx < len(new_row.cells):
                                apply_text_styling_and_borders(new_row.cells[cell_idx], str(val))
                p.text = ""
            else:
                run_replace_mechanism(p, replacements)

        # Process 14-Domain Calculations Matrix Tables
        domain_mapping_keys = [
            ("managementrespons", "management"), ("itsecuritypolic", "policy"), ("humanresource", "human"),
            ("assetmanag", "asset"), ("accesscontrol", "access"), ("cryptograph", "crypto"), ("physical", "physical"),
            ("operation", "operations"), ("communication", "communications"), ("systemacquisition", "development"),
            ("outsourcing", "outsourcing"), ("incident", "incident"), ("businesscontinuity", "continuity"), ("compliance", "compliance")
        ]

        for table in doc.tables:
            if len(table.rows) == 0 or len(table.rows[0].cells) == 0: continue
            first_row_text = "".join(cell.text.lower() for cell in table.rows[0].cells)
            
            for row in table.rows:
                row_header_clean = row.cells[0].text.lower().replace(" ", "").replace(",", "").replace("-", "")
                matched_domain_id = next((internal_id for prefix, internal_id in domain_mapping_keys if prefix in row_header_clean), None)
                
                if matched_domain_id and len(row.cells) >= 5:
                    counts_map = domain_matrix[matched_domain_id]
                    apply_text_styling_and_borders(row.cells[1], str(counts_map["high"]), center=True)
                    apply_text_styling_and_borders(row.cells[2], str(counts_map["medium"]), center=True)
                    apply_text_styling_and_borders(row.cells[3], str(counts_map["low"]), center=True)
                    apply_text_styling_and_borders(row.cells[4], str(counts_map["aoi"]), center=True)
                    continue
                    
                if "grandtotal" in row_header_clean and len(row.cells) >= 5:
                    apply_text_styling_and_borders(row.cells[1], str(global_metrics["high"]), center=True)
                    apply_text_styling_and_borders(row.cells[2], str(global_metrics["medium"]), center=True)
                    apply_text_styling_and_borders(row.cells[3], str(global_metrics["low"]), center=True)
                    apply_text_styling_and_borders(row.cells[4], str(global_metrics["aoi"]), center=True)
                    continue

                for cell in row.cells:
                    for cell_p in cell.paragraphs: run_replace_mechanism(cell_p, replacements)

            # Re-generate Document Sign-Off / Revisions Controls Meta Tables
            if "document revision" in first_row_text or "created by" in first_row_text:
                try:
                    table.cell(0, 0).text = "Version 0.1"
                    table.cell(1, 0).text = "Role"; table.cell(1, 1).text = "Name"; table.cell(1, 2).text = "Action"; table.cell(1, 3).text = "Date"
                    table.cell(2, 0).text = "Created by"; table.cell(2, 1).text = company_name; table.cell(2, 2).text = "Document Revision"; table.cell(2, 3).text = current_date_numeric
                    table.cell(3, 0).text = "Approved by"; table.cell(3, 1).text = ""; table.cell(3, 2).text = "Document Approval"; table.cell(3, 3).text = ""                 
                except Exception: pass
            elif "summary of changes" in first_row_text or "author" in first_row_text:
                try:
                    table.cell(0, 0).text = "Version"; table.cell(0, 1).text = "Date"; table.cell(0, 2).text = "Author"; table.cell(0, 3).text = "Summary of Changes"
                    table.cell(1, 0).text = "0.1"; table.cell(1, 1).text = current_date_numeric; table.cell(1, 2).text = company_name; table.cell(1, 3).text = "Document Creation"
                except Exception: pass

        return_data = io.BytesIO()
        with open(output_path, 'rb') as f:
            return_data.write(f.read())
        return_data.seek(0)
        
        return send_file(
            return_data,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=os.path.basename(output_path)
        )

    except Exception as err:
        return jsonify({"error": f"Internal system automation crash details: {str(err)}"}), 500
        
    finally:
        # Sandbox unlinking garbage collection
        for path in [wab_path, plan_path, code_path, output_path]:
            if os.path.exists(path):
                try: os.unlink(path)
                except Exception: pass

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
